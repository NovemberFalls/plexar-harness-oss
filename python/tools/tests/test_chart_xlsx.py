"""T81 — a real, native Excel chart in the `.xlsx` this package writes.

══ WHY THIS GATE READS THE PACKAGE RATHER THAN THE SOURCE ══════════════════
`chat/office/__init__.py`'s own decision record says a malformed `chart1.xml`
makes the ENTIRE workbook unopenable — a worse failure than no chart at all.
A source scan cannot tell a plausible chart part from a valid one; only a real
reader can, so every workbook this file builds is loaded with `openpyxl`
before anything else is asserted. If `openpyxl` is not installed, the suite
says so loudly rather than silently downgrading to a weaker gate.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from plexar_tools.office.excel._write import xlsx_bytes
from plexar_tools.office._chart import BadChart

try:
    import openpyxl
    _OPENPYXL_VERSION = openpyxl.__version__
except ImportError:  # pragma: no cover - environment fact, not a code path
    openpyxl = None
    _OPENPYXL_VERSION = None


def _line_chart() -> dict:
    return {
        "chart": "line",
        "title": "Revenue",
        "labels": ["Jan", "Feb", "Mar", "Apr"],
        "series": [
            {"name": "Actual", "values": [936204.81, 951002, None, "$972,929"]},
            {"name": "Plan", "values": [900000, 940000, 960000, 970000],
             "color": "#ff0000"},
        ],
        "x_title": "Month",
        "y_title": "USD",
    }


def _bar_chart() -> dict:
    return {
        "chart": "bar",
        "labels": ["North", "South", "East"],
        "series": [{"name": "Units", "values": [10, 20, 30]}],
    }


def test_openpyxl_is_available_or_the_gate_says_so() -> None:
    """**The central gate of this row depends on a real reader existing.**
    If this fails, every other assertion here that "loads clean with
    openpyxl" is unavailable and must not be read as a pass."""
    assert openpyxl is not None, (
        "openpyxl is NOT installed in this environment — the row's central "
        "gate (a real reader loading the generated workbook) could not run. "
        "Every assertion in this file that says 'loads clean with openpyxl' "
        "did not execute.")


def _require_openpyxl():
    if openpyxl is None:
        pytest.skip("openpyxl not installed — see "
                    "test_openpyxl_is_available_or_the_gate_says_so")


# ══ THE CENTRAL GATE: A REAL READER MUST OPEN EVERY WORKBOOK THIS ROW MAKES ═

def test_a_line_chart_workbook_loads_clean_with_openpyxl() -> None:
    _require_openpyxl()
    data = xlsx_bytes([["header"]], sheet="P&L", charts=[_line_chart()])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert len(ws._charts) == 1
    assert ws._charts[0].__class__.__name__ == "LineChart"


def test_a_bar_chart_workbook_loads_clean_with_openpyxl() -> None:
    _require_openpyxl()
    data = xlsx_bytes([], sheet="Regions", charts=[_bar_chart()])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert len(ws._charts) == 1
    assert ws._charts[0].__class__.__name__ == "BarChart"


def test_two_charts_on_one_sheet_both_load() -> None:
    _require_openpyxl()
    data = xlsx_bytes([], sheet="Both", charts=[_line_chart(), _bar_chart()])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert len(ws._charts) == 2


def test_no_charts_argument_is_unaffected_and_still_loads() -> None:
    """**Existing callers pass no `charts` kwarg at all.** Backward
    compatibility is asserted, not assumed."""
    _require_openpyxl()
    data = xlsx_bytes([["a", "b"], ["1", "2"]], sheet="Plain")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert len(ws._charts) == 0
    assert ws["A1"].value == "a"


# ══ THE NUMERIC-CELL NARROWING ══════════════════════════════════════════════

def test_charted_cells_are_numeric_and_every_other_cell_stays_inline_string() -> None:
    """**The narrowing must be scoped to exactly the charted range.** A
    caller-supplied row of strings must never be reinterpreted as numbers,
    and a chart's own label column must stay text."""
    data = xlsx_bytes([["Region", "Q1"], ["North", "10"]],
                       sheet="Mix", charts=[_bar_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")

    # The caller's own rows: still inlineStr, untouched.
    assert '<c r="A1" t="inlineStr"><is><t xml:space="preserve">Region' in sheet_xml
    assert '<c r="B2" t="inlineStr"><is><t xml:space="preserve">10' in sheet_xml

    # The chart's header row (series names) is text.
    assert '<t xml:space="preserve">Units</t>' in sheet_xml

    # The chart's label column (categories, column A) stays text: every A-cell
    # from row 3 down (blank separator, then header, then 3 label rows) must
    # have no `<v>`.
    a_cells_after_caller_rows = [
        seg for seg in sheet_xml.split('<c r="A')[3:]  # A1, A2 are the caller's
    ]
    assert a_cells_after_caller_rows, "expected label cells in column A"
    for seg in a_cells_after_caller_rows:
        cell = seg.split("</c>")[0]
        assert "<v>" not in cell, f"a chart label cell must not be numeric: {cell!r}"

    # The chart's VALUE cells (column B, below its header) are numeric.
    b_segments = [seg.split("</c>")[0] for seg in sheet_xml.split('<c r="B')[3:]]
    numeric_b = [seg for seg in b_segments if "<v>" in seg]
    assert len(numeric_b) == 3
    for cell in numeric_b:
        assert "inlineStr" not in cell
    assert "<v>10</v>" in sheet_xml and "<v>20</v>" in sheet_xml and "<v>30</v>" in sheet_xml


def test_string_numbers_from_the_model_are_parsed_into_real_numeric_cells() -> None:
    """`_chart.parse_chart` accepts `"$972,929"` and turns it into a float —
    that float must reach the sheet as `<v>`, not as the original string."""
    data = xlsx_bytes([], sheet="Parsed", charts=[_line_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "$972,929" not in sheet_xml
    assert "<v>972929</v>" in sheet_xml


def test_a_gap_value_is_a_blank_cell_not_a_zero() -> None:
    """`None` in a series is a gap (INVARIANT-ENVELOPE) — it must be OMITTED,
    never written as `<v>0</v>`."""
    data = xlsx_bytes([], sheet="Gap", charts=[_line_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    # Row 6 (start_row(4)+1(header)+1..) holds March, whose Actual value is
    # None. Column B is Actual; that cell must not exist as a `<v>` at all.
    # We assert structurally: the number of <v> cells in the data grid for
    # column B is 3 (Jan, Feb, Apr), not 4.
    b_values = [seg for seg in sheet_xml.split("<c r=\"B")
                if seg and seg[0].isdigit()]
    numeric_b = [seg for seg in b_values if "<v>" in seg.split("</c>")[0]]
    assert len(numeric_b) == 3


# ══ PACKAGE CLOSURE — EVERY RELS TARGET RESOLVES ═══════════════════════════

def test_package_is_closed_every_rels_target_resolves() -> None:
    data = xlsx_bytes([], sheet="Closed", charts=[_line_chart(), _bar_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = set(zf.namelist())

    import re
    for rels_name in [n for n in names if n.endswith(".rels")]:
        base = rels_name.rsplit("_rels/", 1)[0]
        xml = zf.read(rels_name).decode("utf-8")
        for target in re.findall(r'Target="([^"]+)"', xml):
            resolved = _resolve(base, target)
            assert resolved in names, (
                f"{rels_name} points at {target!r} -> {resolved!r}, "
                f"which is not in the package")


def _resolve(base: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    parts = base.rstrip("/").split("/") if base.rstrip("/") else []
    for piece in target.split("/"):
        if piece == "..":
            parts.pop()
        elif piece == ".":
            continue
        else:
            parts.append(piece)
    return "/".join(parts)


def test_content_types_overrides_exist_for_chart_and_drawing_parts() -> None:
    data = xlsx_bytes([], sheet="CT", charts=[_line_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    ct = zf.read("[Content_Types].xml").decode("utf-8")
    assert '/xl/charts/chart1.xml' in ct
    assert '/xl/drawings/drawing1.xml' in ct
    assert "drawingml.chart+xml" in ct
    assert "officedocument.drawing+xml" in ct


def test_chart_formula_references_the_sanitised_sheet_name() -> None:
    """A sheet named with characters Excel forbids is sanitised at ~L152;
    the chart's `c:f` formulas must use THAT name, not the caller's input."""
    data = xlsx_bytes([], sheet="Q1/Q2:Report*", charts=[_bar_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    chart_xml = zf.read("xl/charts/chart1.xml").decode("utf-8")
    assert "Q1/Q2:Report*" not in chart_xml
    assert "Q1-Q2-Report-" in chart_xml


def test_styles_xml_now_has_a_relationship_from_workbook_rels() -> None:
    """Pre-existing bug fixed in this row: `xl/styles.xml` shipped with no
    relationship naming it from `xl/_rels/workbook.xml.rels`."""
    data = xlsx_bytes([["a"]], sheet="S")
    zf = zipfile.ZipFile(io.BytesIO(data))
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    assert "styles.xml" in rels


# ══ WATCH-TO-FAIL: THE READER-LOAD ASSERTION MUST BE ABLE TO FAIL ══════════

def test_a_corrupted_chart_part_is_rejected_by_the_real_reader() -> None:
    """**Proves the gate has teeth.** `chart1.xml` is corrupted IN MEMORY —
    never written to disk, so nothing on disk is ever wrong (T5's rule,
    inherited) — and the same reader used by every arm above must refuse it.
    """
    _require_openpyxl()
    data = xlsx_bytes([], sheet="Sabotaged", charts=[_bar_chart()])
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = zf.namelist()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name in names:
            raw = zf.read(name)
            if name == "xl/charts/chart1.xml":
                # Truncate mid-element: no closing tags at all.
                raw = raw[: len(raw) // 3]
            out.writestr(name, raw)
    corrupted = buf.getvalue()

    # Confirm the sabotage actually took (R23): the corrupted bytes are not
    # the well-formed original.
    assert corrupted != data
    with pytest.raises(Exception):
        wb = openpyxl.load_workbook(io.BytesIO(corrupted))
        # Some lenient parsers open the workbook but the chart access itself
        # raises, or produces no chart at all where one was expected — either
        # is a real reader detecting the corruption, so both are accepted as
        # the failure this arm exists to prove is reachable.
        ws = wb.active
        assert len(ws._charts) == 1


def test_a_bad_chart_block_raises_before_any_xml_is_built() -> None:
    """`parse_chart` is the validated boundary — a malformed block must never
    reach the OOXML writer at all."""
    with pytest.raises(BadChart):
        xlsx_bytes([], sheet="Bad", charts=[{"chart": "pie", "labels": ["a", "b"],
                                              "series": [{"values": [1, 2]}]}])
